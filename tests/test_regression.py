"""
Tests de non-régression — BotShortTrade

Usage :
    python3 -m unittest discover -s tests -v

Aucun appel réseau, aucun ordre réel, aucun message Discord :
  - dotenv neutralisé (le vrai .env n'est jamais lu) + clé privée factice
  - SDK Hyperliquid remplacé par des stubs avant l'import de main.py
  - requests.get/post lèvent une erreur si un test oublie de les mocker
  - exécution dans un dossier temporaire (trades.xlsx, bot_state.json… factices)
"""
import json, os, sys, tempfile, types, unittest
from unittest import mock

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# ── Isolation AVANT l'import de main.py ──────────────────────────────────────
_dotenv = types.ModuleType("dotenv")
_dotenv.load_dotenv = lambda *a, **k: None
sys.modules["dotenv"] = _dotenv

os.environ.update({
    "HL_PRIVATE_KEY": "0x" + "11" * 32,   # clé factice
    "HL_TESTNET": "1",
    "DISCORD_WEBHOOK_URL": "",
    "LEVERAGE": "5", "CAPITAL_PCT": "0.15", "VOL_MULTIPLIER": "1.2",
    "MIN_MOVE_PCT": "0.0002", "MAX_TRADES": "3", "HEARTBEAT_HOURS": "6",
})

class _StubInfo:
    def __init__(self, *a, **k): pass
class _StubExchange:
    def __init__(self, *a, **k): pass
for name, attr, cls in [("hyperliquid.info", "Info", _StubInfo),
                        ("hyperliquid.exchange", "Exchange", _StubExchange)]:
    mod = types.ModuleType(name); setattr(mod, attr, cls); sys.modules[name] = mod
if "hyperliquid" not in sys.modules:
    sys.modules["hyperliquid"] = types.ModuleType("hyperliquid")
try:
    import hyperliquid.utils.constants  # noqa: F401  (vrai module, sans réseau)
except Exception:
    _u = types.ModuleType("hyperliquid.utils"); _c = types.ModuleType("hyperliquid.utils.constants")
    _c.TESTNET_API_URL = _c.MAINNET_API_URL = "http://stub"; _u.constants = _c
    sys.modules["hyperliquid.utils"] = _u; sys.modules["hyperliquid.utils.constants"] = _c

import requests
def _no_network(*a, **k):
    raise AssertionError(f"Appel réseau interdit dans les tests : {a[:1]}")
requests.get = requests.post = _no_network

_TMP = tempfile.mkdtemp(prefix="bst_tests_")
os.chdir(_TMP)
sys.path.insert(0, ROOT)
import contextlib, io
with contextlib.redirect_stdout(io.StringIO()):
    import main as m
import openpyxl


# ── Helpers ──────────────────────────────────────────────────────────────────
def _ok(oid=2):
    return {"status": "ok", "response": {"data": {"statuses": [{"resting": {"oid": oid}}]}}}

def _err(msg="Invalid TP/SL price"):
    return {"status": "ok", "response": {"data": {"statuses": [{"error": msg}]}}}

class FakeExchange:
    """Simule Hyperliquid : fill à avgPx, réponses SL/TP/close configurables."""
    def __init__(self, avg_px=2010.0, sl=None, tp=None, close_ok=True):
        self.avg_px, self.close_ok = avg_px, close_ok
        self.sl = sl or [_ok(11)]
        self.tp = tp or [_ok(12)]
        self.orders, self.closed = [], 0
    def update_leverage(self, *a, **k): return {"status": "ok"}
    def market_open(self, *a, **k):
        return {"status": "ok", "response": {"data": {"statuses": [
            {"filled": {"oid": 1, "avgPx": str(self.avg_px), "totalSz": "0.04"}}]}}}
    def order(self, coin, is_buy, qty, limit_px, order_type, reduce_only):
        trig = order_type["trigger"]; self.orders.append(trig)
        seq = self.sl if trig["tpsl"] == "sl" else self.tp
        return seq.pop(0) if len(seq) > 1 else seq[0]
    def market_close(self, coin, slippage=0.05):
        self.closed += 1
        return {"status": "ok"} if self.close_ok else {"status": "err"}
    def cancel(self, *a, **k): return {"status": "ok"}

class FakeInfo:
    def __init__(self, mids=None, perp=None, spot=None, fills=None):
        self.mids = mids or {"ETH": "2000"}
        self.perp, self.spot, self.fills = perp or {}, spot or {"balances": []}, fills or []
    def all_mids(self): return self.mids
    def user_state(self, addr): return self.perp
    def spot_user_state(self, addr): return self.spot
    def user_fills(self, addr): return self.fills

SIG = {"symbol": "ETH", "signal": "buy", "score": 0.8, "atr_m5": 5.0, "price": 2000,
       "conditions": {"mode": "MR", "regime": "RANGING", "rsi": 50, "bbp": 0.3, "vol_ratio": 1.5}}


class BaseTest(unittest.TestCase):
    def setUp(self):
        for f in (m.EXCEL_FILE, m.STATE_FILE, m.LESSONS_FILE):
            if os.path.exists(f):
                os.remove(f)
        self.alerts = []
        self._p = [
            mock.patch.object(m, "send_discord_alert",
                              lambda t, d, color=0: self.alerts.append((t, d))),
            mock.patch.object(m.time, "sleep", lambda s: None),
            mock.patch("sys.stdout", new_callable=io.StringIO),
        ]
        for p in self._p: p.start()
    def tearDown(self):
        for p in self._p: p.stop()
    def titles(self):
        return [t for t, _ in self.alerts]


# ── place_order : SL/TP et sécurité stop-loss ────────────────────────────────
class TestPlaceOrder(BaseTest):
    def run_order(self, ex, sig=SIG):
        with mock.patch.object(m, "exchange", ex), mock.patch.object(m, "info", FakeInfo()):
            wb = m.init_excel()
            return m.place_order(dict(sig), wb), wb

    def test_sl_tp_calcules_sur_prix_execute(self):
        pos, _ = self.run_order(FakeExchange(avg_px=2010.0))
        self.assertEqual(pos["entry"], 2010.0)
        self.assertAlmostEqual(pos["sl"], 2010.0 - m.ATR_SL_MULT * 5.0)
        self.assertAlmostEqual(pos["tp"], 2010.0 + m.ATR_TP_MULT * 5.0)
        self.assertEqual(pos["sl_oid"], 11)

    def test_sl_tp_short(self):
        sig = dict(SIG, signal="sell")
        pos, _ = self.run_order(FakeExchange(avg_px=1990.0), sig)
        self.assertAlmostEqual(pos["sl"], 1990.0 + m.ATR_SL_MULT * 5.0)
        self.assertAlmostEqual(pos["tp"], 1990.0 - m.ATR_TP_MULT * 5.0)

    def test_sl_refuse_puis_accepte_au_retry(self):
        ex = FakeExchange(sl=[_err(), _ok(33)])
        pos, _ = self.run_order(ex)
        self.assertEqual(pos["sl_oid"], 33)
        self.assertEqual(ex.closed, 0)

    def test_sl_refuse_deux_fois_ferme_la_position(self):
        ex = FakeExchange(sl=[_err()])
        with self.assertRaises(RuntimeError):
            self.run_order(ex)
        self.assertEqual(ex.closed, 1)
        self.assertTrue(any("SL refusé" in t for t in self.titles()))
        self.assertFalse(any(o["tpsl"] == "tp" for o in ex.orders))   # pas de TP posé

    def test_sl_status_error_global(self):
        ex = FakeExchange(sl=[{"status": "err", "response": "rate limited"}])
        with self.assertRaises(RuntimeError):
            self.run_order(ex)
        self.assertEqual(ex.closed, 1)

    def test_sl_refuse_et_fermeture_echoue_alerte_critique(self):
        ex = FakeExchange(sl=[_err()], close_ok=False)
        pos, _ = self.run_order(ex)          # position renvoyée pour rester suivie
        self.assertIsNone(pos["sl_oid"])
        self.assertTrue(any("SANS STOP-LOSS" in t for t in self.titles()))

    def test_tp_refuse_alerte_mais_position_gardee(self):
        ex = FakeExchange(tp=[_err()])
        pos, _ = self.run_order(ex)
        self.assertEqual(pos["sl_oid"], 11)
        self.assertEqual(ex.closed, 0)
        self.assertTrue(any("TP refusé" in t for t in self.titles()))

    def test_ligne_excel_mode_regime_score(self):
        _, wb = self.run_order(FakeExchange())
        ws = wb["Trades"]
        row = [c.value for c in ws[ws.max_row]]
        hdr = [c.value for c in ws[1]]
        self.assertEqual(row[hdr.index("Mode")], "MR")
        self.assertEqual(row[hdr.index("Régime")], "RANGING")
        self.assertEqual(row[hdr.index("Score")], 0.8)
        self.assertEqual(hdr[14], "P&L final $")   # index utilisé par analyze_weekly.py


# ── Excel ────────────────────────────────────────────────────────────────────
class TestExcel(BaseTest):
    def test_migration_ancien_fichier_15_colonnes(self):
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Trades"
        ws.append(m.COL_HEADERS[:15]); ws.append(["1", "2026-01-01 00:00"] + [0] * 13)
        wb.save(m.EXCEL_FILE)
        ws2 = m.init_excel()["Trades"]
        self.assertEqual([c.value for c in ws2[1]], m.COL_HEADERS)
        self.assertEqual(ws2.cell(2, 1).value, "1")   # données conservées


# ── Équité ───────────────────────────────────────────────────────────────────
class TestEquity(BaseTest):
    def eq(self, perp, spot):
        with mock.patch.object(m, "info", FakeInfo(spot=spot)):
            return m.get_equity(perp)

    def test_unified_marge_pas_comptee_deux_fois(self):
        # Valeurs réelles relevées le 2026-10-16 avec 1 position DOGE ouverte
        perp = {"withdrawable": "0.0", "crossMarginSummary":
                {"accountValue": "16.556911", "totalMarginUsed": "16.68719"}}
        spot = {"balances": [{"coin": "USDC", "total": "111.343455", "hold": "16.68719"}]}
        self.assertAlmostEqual(self.eq(perp, spot), 111.21, places=2)

    def test_unified_sans_position(self):
        perp = {"withdrawable": "0.0", "crossMarginSummary":
                {"accountValue": "0.0", "totalMarginUsed": "0.0"}}
        spot = {"balances": [{"coin": "USDC", "total": "111.5483", "hold": "0.0"}]}
        self.assertAlmostEqual(self.eq(perp, spot), 111.5483, places=4)

    def test_compte_classique_perp_seul(self):
        perp = {"withdrawable": "80", "crossMarginSummary":
                {"accountValue": "100", "totalMarginUsed": "20"}}
        self.assertAlmostEqual(self.eq(perp, {"balances": []}), 100.0)

    def test_compte_classique_perp_plus_spot(self):
        perp = {"crossMarginSummary": {"accountValue": "100", "totalMarginUsed": "20"}}
        spot = {"balances": [{"coin": "USDC", "total": "50", "hold": "0"}]}
        self.assertAlmostEqual(self.eq(perp, spot), 150.0)

    def test_spot_indisponible_leve_au_lieu_de_zero(self):
        """25/10 : 504 sur spot_user_state → équité lue $0.00 → faux flash crash -100 %."""
        info = FakeInfo(spot={})
        info.spot_user_state = mock.Mock(side_effect=Exception("(504, 'Gateway Timeout')"))
        perp = {"crossMarginSummary": {"accountValue": "0.0", "totalMarginUsed": "0.0"}}
        with mock.patch.object(m, "info", info), self.assertRaises(Exception):
            m.get_equity(perp)

    def test_user_state_vide_leve(self):
        with mock.patch.object(m, "info", FakeInfo(spot={"balances": []})), self.assertRaises(ValueError):
            m.get_equity({})


# ── P&L : frais d'ouverture ──────────────────────────────────────────────────
class TestSharedWallet(BaseTest):
    POS = [{"position": {"coin": "ETH", "szi": "0.01", "entryPx": "2000", "unrealizedPnl": "1.5"}},
           {"position": {"coin": "SOL", "szi": "-0.5", "entryPx": "100", "unrealizedPnl": "-3"}},
           {"position": {"coin": "XRP", "szi": "0", "entryPx": "1", "unrealizedPnl": "0"}}]

    def test_account_coins(self):
        self.assertEqual(m.account_coins({"assetPositions": self.POS}), {"ETH", "SOL"})
        self.assertEqual(m.account_coins(None), set())

    def test_pnl_latent_limite_a_nos_positions(self):
        st = {"assetPositions": self.POS}
        self.assertAlmostEqual(m.get_unrealized_pnl(st, {"ETH"}), 1.5)
        self.assertAlmostEqual(m.get_unrealized_pnl(st), -1.5)

    def test_reprise_n_adopte_que_nos_positions(self):
        m.save_positions_state([{"id": "o1", "symbol": "ETH", "side": "BUY", "entry": 2000,
                                 "qty": 0.01, "sl": 1990, "tp": 2030, "sl_oid": 7, "atr": 5,
                                 "open_fee": 0.01, "opened_at": 1.0, "conditions": {}}])
        with mock.patch.object(m, "info", FakeInfo(perp={"assetPositions": self.POS})):
            rec, _ = m.recover_open_positions()
        self.assertEqual([p["symbol"] for p in rec], ["ETH"])
        self.assertEqual(rec[0]["sl_oid"], 7)
        self.assertEqual(self.titles(), ["⚠️ Position non reprise"])


class TestFees(BaseTest):
    def test_zombie_pnl_deduit_frais_ouverture_estimes(self):
        fills = [{"coin": "ETH", "dir": "Close Long", "time": 1, "fee": "0.04", "closedPnl": "1.00"}]
        pos = {"id": "1", "symbol": "ETH", "side": "BUY", "entry": 2000, "qty": 0.04,
               "sl_oid": None, "open_fee": 0.0, "conditions": {}}
        with mock.patch.object(m, "exchange", FakeExchange()), \
             mock.patch.object(m, "info", FakeInfo(fills=fills)), \
             mock.patch.object(m, "update_order_status", lambda *a, **k: None), \
             mock.patch.object(m, "record_lesson", lambda *a, **k: None):
            pnl = m.force_close_zombie(m.init_excel(), pos)
        self.assertAlmostEqual(pnl, 1.00 - 0.04 - 0.04, places=2)


    # Fills réels du trade DOGE du 2026-10-16 : SL exécuté en 3 fills partiels.
    # Avant correctif : seul le dernier fill était compté → P&L −0.08 au lieu de −0.37.
    DOGE_FILLS = [
        {"coin": "DOGE", "dir": "Open Long",  "time": 1792136736618, "oid": 576377349716,
         "sz": "1016.0", "closedPnl": "0.0", "fee": "0.036116"},
        {"coin": "DOGE", "dir": "Close Long", "time": 1792138739787, "oid": 576377359631,
         "sz": "222.0", "closedPnl": "-0.0666", "fee": "0.007862"},
        {"coin": "DOGE", "dir": "Close Long", "time": 1792138739787, "oid": 576377359631,
         "sz": "677.0", "closedPnl": "-0.195653", "fee": "0.023981"},
        {"coin": "DOGE", "dir": "Close Long", "time": 1792138739787, "oid": 576377359631,
         "sz": "117.0", "closedPnl": "-0.033579", "fee": "0.004144"},
        # Ancien trade DOGE fermé plus tôt (autre oid) : ne doit PAS être additionné
        {"coin": "DOGE", "dir": "Close Short", "time": 1792000000000, "oid": 111,
         "sz": "500.0", "closedPnl": "5.0", "fee": "0.02"},
    ]
    DOGE_POS = {"id": "576377349716", "symbol": "DOGE", "side": "BUY", "entry": 0.082287,
                "qty": 1016.0, "sl": 0.082014, "tp": 0.083016, "sl_oid": 576377359631,
                "atr": 0.000182, "open_fee": 0.0, "conditions": {"mode": "MR"}}

    def test_sl_en_fills_partiels_tous_additionnes(self):
        rec = {}
        with mock.patch.object(m, "info", FakeInfo(fills=self.DOGE_FILLS)), \
             mock.patch.object(m, "update_order_status",
                               lambda wb, oid, status, pnl=None: rec.update(status=status, pnl=pnl)), \
             mock.patch.object(m, "record_lesson", lambda *a, **k: None):
            result = m.check_position(None, dict(self.DOGE_POS), cached_state={"assetPositions": []})
        self.assertEqual(result, "SL touché")
        # gross −0.295832 − frais ouverture réels 0.036116 − frais sortie 0.035987
        self.assertAlmostEqual(rec["pnl"], -0.37, places=2)

    def test_closing_pnl_frais_ouverture_reels(self):
        gross, open_fee, close_fee = m._closing_pnl(self.DOGE_POS, self.DOGE_FILLS)
        self.assertAlmostEqual(gross, -0.295832, places=6)
        self.assertAlmostEqual(open_fee, 0.036116, places=6)
        self.assertAlmostEqual(close_fee, 0.035987, places=6)


# ── Revue 2026-10-16 : zombie, attribution des fermetures, surveillance SL ───
class OrdersInfo(FakeInfo):
    def __init__(self, orders=None, **k):
        super().__init__(**k); self.orders = orders or []
    def frontend_open_orders(self, addr): return self.orders

class LogExchange(FakeExchange):
    def __init__(self, close_resp=None, **k):
        super().__init__(**k); self.log, self.close_resp = [], close_resp
    def market_close(self, coin, slippage=0.05):
        self.log.append("close"); return self.close_resp
    def cancel(self, coin, oid):
        self.log.append(("cancel", oid)); return {"status": "ok"}
    def order(self, *a, **k):
        self.log.append("order"); return super().order(*a, **k)

class TestZombie(BaseTest):
    POS = {"id": "1", "symbol": "ETH", "side": "BUY", "entry": 2000, "qty": 0.04,
           "sl_oid": 11, "open_fee": 0.0, "conditions": {}}

    def close(self, resp):
        ex = LogExchange(close_resp=resp)
        with mock.patch.object(m, "exchange", ex), mock.patch.object(m, "info", FakeInfo()), \
             mock.patch.object(m, "update_order_status", lambda *a, **k: None), \
             mock.patch.object(m, "record_lesson", lambda *a, **k: None):
            return m.force_close_zombie(None, dict(self.POS)), ex

    def test_fermeture_puis_annulation_sl(self):
        pnl, ex = self.close({"status": "ok", "response": {"data": {"statuses": [{"filled": {}}]}}})
        self.assertIsNotNone(pnl)
        self.assertEqual(ex.log, ["close", ("cancel", 11)])

    def test_fermeture_echouee_sl_conserve(self):
        pnl, ex = self.close({"status": "err", "response": "insufficient liquidity"})
        self.assertIsNone(pnl)
        self.assertEqual(ex.log, ["close"])                  # SL jamais annulé
        self.assertTrue(any("fermeture échouée" in t for t in self.titles()))

    def test_fermeture_statut_erreur_sl_conserve(self):
        pnl, ex = self.close({"status": "ok", "response": {"data": {"statuses": [{"error": "x"}]}}})
        self.assertIsNone(pnl)
        self.assertEqual(ex.log, ["close"])


class TestClosingAttribution(BaseTest):
    def test_fermeture_anterieure_a_l_ouverture_ignoree(self):
        fills = [{"coin": "DOGE", "dir": "Close Short", "time": 1_000_000, "oid": 9,
                  "closedPnl": "5.0", "fee": "0.02"}]
        pos = {"id": "1", "symbol": "DOGE", "opened_at": 2_000.0}   # ouvert après ce fill
        self.assertIsNone(m._closing_pnl(pos, fills))

    def test_position_absente_sans_fill_attend_3_cycles(self):
        pos = {"id": "1", "symbol": "DOGE", "side": "BUY", "entry": 0.08, "qty": 1000,
               "opened_at": 2_000_000_000.0, "conditions": {}}
        rec = {}
        with mock.patch.object(m, "info", FakeInfo(fills=[])), \
             mock.patch.object(m, "update_order_status",
                               lambda wb, oid, status, pnl=None: rec.update(status=status)), \
             mock.patch.object(m, "record_lesson", lambda *a, **k: None):
            r1 = m.check_position(None, pos, cached_state={"assetPositions": []})
            r2 = m.check_position(None, pos, cached_state={"assetPositions": []})
            r3 = m.check_position(None, pos, cached_state={"assetPositions": []})
        self.assertEqual((r1, r2), (None, None))
        self.assertEqual(r3, "Fermé (inconnu)")
        self.assertEqual(rec["status"], "Fermé (inconnu)")


class TestSLWatchdog(BaseTest):
    POS = {"id": "1", "symbol": "ETH", "side": "BUY", "entry": 2000.0, "qty": 0.04,
           "sl": 1992.5, "sl_oid": None}
    STOP = {"coin": "ETH", "isTrigger": True, "reduceOnly": True, "orderType": "Stop Market", "oid": 55}
    TP   = {"coin": "ETH", "isTrigger": True, "reduceOnly": True, "orderType": "Take Profit Market", "oid": 56}

    def watch(self, orders, ex=None):
        pos, ex = dict(self.POS), ex or LogExchange()
        with mock.patch.object(m, "info", OrdersInfo(orders=orders)), mock.patch.object(m, "exchange", ex):
            n = m.ensure_sl_orders([pos])
        return n, pos, ex

    def test_sl_present_resynchronise_oid(self):
        n, pos, ex = self.watch([self.STOP, self.TP])
        self.assertEqual(n, 0)
        self.assertEqual(pos["sl_oid"], 55)
        self.assertEqual(ex.log, [])

    def test_sl_absent_repose(self):
        n, pos, ex = self.watch([self.TP])                 # seul le TP existe
        self.assertEqual(n, 1)
        self.assertEqual(ex.orders[0]["triggerPx"], 1992.5)
        self.assertEqual(ex.orders[0]["tpsl"], "sl")
        self.assertEqual(pos["sl_oid"], 11)
        self.assertTrue(any("SL manquant" in t for t in self.titles()))

    def test_sl_absent_et_refuse_alerte_critique(self):
        n, pos, ex = self.watch([], LogExchange(sl=[_err()]))
        self.assertEqual(n, 0)
        self.assertTrue(any("SANS STOP-LOSS" in t for t in self.titles()))


class TestDailyPnl(BaseTest):
    def test_pnl_du_jour_par_date_de_fermeture(self):
        wb = m.init_excel(); ws = wb["Trades"]
        today = m.datetime.now(m.timezone.utc).strftime("%Y-%m-%d")
        base = ["x"] * 13
        ws.append(["1", "2000-01-01 23:50"] + base[:11] + ["SL touché", -1.0, "MR", "RANGING", 0.8, f"{today} 00:10"])
        ws.append(["2", f"{today} 08:00"] + base[:11] + ["TP touché", 0.5])          # ancienne ligne sans date
        ws.append(["3", f"{today} 09:00"] + base[:11] + ["SL touché", -0.3, "MR", "RANGING", 0.8, "2000-01-02 00:00"])
        self.assertAlmostEqual(m.calc_daily_pnl(wb), -0.5)

    def test_date_fermeture_ecrite(self):
        with mock.patch.object(m, "exchange", FakeExchange()), mock.patch.object(m, "info", FakeInfo()):
            wb = m.init_excel(); m.place_order(dict(SIG), wb)
        oid = wb["Trades"].cell(2, 1).value
        m.update_order_status(wb, oid, "TP touché", 0.7)
        self.assertTrue(str(wb["Trades"].cell(2, 19).value).startswith(
            m.datetime.now(m.timezone.utc).strftime("%Y-%m-%d")))


# ── Trailing stop ────────────────────────────────────────────────────────────
class TrailExchange(FakeExchange):
    def __init__(self, resp):
        super().__init__(); self.resp, self.log = resp, []
    def order(self, *a, **k):
        self.log.append("order"); return self.resp
    def cancel(self, coin, oid):
        self.log.append(("cancel", oid)); return {"status": "ok"}

class TestTrailing(BaseTest):
    POS = {"symbol": "ETH", "side": "BUY", "entry": 2000.0, "qty": 0.04,
           "sl": 1992.5, "sl_oid": 11, "atr": 5.0}

    def trail(self, resp, price=2010.0):   # +2×ATR → phase trailing
        pos, ex = dict(self.POS), TrailExchange(resp)
        with mock.patch.object(m, "exchange", ex):
            m.trail_sl(pos, price)
        return pos, ex

    def test_nouveau_sl_pose_avant_annulation_ancien(self):
        pos, ex = self.trail(_ok(77))
        self.assertEqual(ex.log, ["order", ("cancel", 11)])
        self.assertEqual(pos["sl_oid"], 77)
        self.assertAlmostEqual(pos["sl"], 2010.0 - m.TRAIL_DIST * 5.0)
        self.assertTrue(pos["trailed"])

    def test_nouveau_sl_refuse_ancien_conserve(self):
        pos, ex = self.trail(_err())
        self.assertEqual(ex.log, ["order"])            # ancien SL jamais annulé
        self.assertEqual(pos["sl_oid"], 11)
        self.assertEqual(pos["sl"], 1992.5)
        self.assertNotIn("trailed", pos)

    def test_nouveau_sl_status_erreur_ancien_conserve(self):
        pos, ex = self.trail({"status": "err", "response": "rate limited"})
        self.assertEqual(ex.log, ["order"])
        self.assertEqual(pos["sl_oid"], 11)

    def test_breakeven_couvre_les_frais(self):
        pos, ex = self.trail(_ok(77), price=2006.0)   # +1.2×ATR → phase breakeven
        self.assertAlmostEqual(pos["sl"], 2000.0 * (1 + m.BE_BUFFER_PCT / 100))

    def test_breakeven_short(self):
        pos = dict(self.POS, side="SELL", sl=2007.5)
        with mock.patch.object(m, "exchange", TrailExchange(_ok(78))):
            m.trail_sl(pos, 1994.0)
        self.assertAlmostEqual(pos["sl"], 2000.0 * (1 - m.BE_BUFFER_PCT / 100))

    def test_pas_de_trail_sous_1_atr(self):
        pos, ex = self.trail(_ok(77), price=2003.0)
        self.assertEqual(ex.log, [])


# ── Apprentissage ────────────────────────────────────────────────────────────
class TestLessons(BaseTest):
    POS = {"symbol": "DOGE", "side": "BUY",
           "conditions": {"rsi": 52.9, "bbp": 0.346, "regime": "RANGING", "mode": "MR"}}

    def setUp(self):
        super().setUp()
        m._lessons_cache.clear(); m._lessons_cache_ts = 0.0

    def test_trail_ferme_a_perte_compte_comme_perte(self):
        m.record_lesson(self.POS, "TRAIL touché", -0.05)
        b = m._load_lessons()["buckets"]["coin_DOGE|buy"]
        self.assertEqual((b["w"], b["l"]), (0, 1))

    def test_trail_ferme_en_gain_compte_comme_gain(self):
        m.record_lesson(self.POS, "TRAIL touché", 0.40)
        b = m._load_lessons()["buckets"]["coin_DOGE|buy"]
        self.assertEqual((b["w"], b["l"]), (1, 0))

    def pen(self):
        return m.get_penalty("DOGE", "BUY", 52.9, 0.346, "RANGING", "MR")

    def test_penalite_apres_6_pertes_avec_oubli(self):
        for _ in range(5):
            m.record_lesson(self.POS, "SL touché", -0.37)
        self.assertEqual(self.pen(), 0.0)   # oubli : 5 pertes ≈ 4.71 effectives < 5
        m.record_lesson(self.POS, "SL touché", -0.37)
        self.assertAlmostEqual(self.pen(), 0.25)   # 4 buckets à 0 % WR → plafond

    def test_bonus_si_historique_gagnant(self):
        for _ in range(8):
            m.record_lesson(self.POS, "TP touché", 0.70)
        self.assertAlmostEqual(self.pen(), -0.15)  # bonus plafonné à 0.15

    def test_oubli_progressif(self):
        m.record_lesson(self.POS, "SL touché", -0.37)
        m.record_lesson(self.POS, "TP touché", 0.70)
        b = m._load_lessons()["buckets"]["coin_DOGE|buy"]
        self.assertAlmostEqual(b["l"], 0.97)
        self.assertAlmostEqual(b["w"], 1.0)

    def test_fantomes_comptent_a_moitie(self):
        for _ in range(10):
            m.record_lesson(self.POS, "SL", -0.5, ghost=True)
        self.assertEqual(self.pen(), 0.0)          # 10 fantômes ≈ 4.3 effectifs
        for _ in range(2):
            m.record_lesson(self.POS, "SL", -0.5, ghost=True)
        self.assertGreater(self.pen(), 0.0)
        b = m._load_lessons()["buckets"]["coin_DOGE|buy"]
        self.assertEqual((b["w"], b["l"]), (0, 0))  # compteurs réels intacts
        self.assertAlmostEqual(b["gpnl_pct"], -6.0)


# ── Mode fantôme ─────────────────────────────────────────────────────────────
def _m1(ts, o, h, l, c):
    return {"ts": ts, "o": o, "h": h, "l": l, "c": c, "v": 1.0}

class TestGhosts(BaseTest):
    def setUp(self):
        super().setUp()
        if os.path.exists(m.GHOST_FILE):
            os.remove(m.GHOST_FILE)
        m._ghosts = None
        m._lessons_cache.clear(); m._lessons_cache_ts = 0.0
        self.wb = m.init_excel()

    def g(self, side="buy"):
        return {"side": side, "entry": 100.0, "sl": 99.0 if side == "buy" else 101.0,
                "tp": 102.0 if side == "buy" else 98.0, "checked_ts": 0, "opened_at": 1000.0}

    def test_tp_touche(self):
        g = self.g()
        self.assertEqual(m._ghost_resolve(g, [_m1(1, 100, 101, 99.5, 101), _m1(2, 101, 102.1, 100.5, 102)], 1100),
                         ("TP", 102.0))
        self.assertEqual(g["checked_ts"], 2)

    def test_sl_et_tp_meme_bougie_compte_sl(self):
        self.assertEqual(m._ghost_resolve(self.g(), [_m1(1, 100, 102.5, 98.5, 100)], 1100), ("SL", 99.0))

    def test_short(self):
        self.assertEqual(m._ghost_resolve(self.g("sell"), [_m1(1, 100, 100.2, 97.9, 98)], 1100), ("TP", 98.0))
        self.assertEqual(m._ghost_resolve(self.g("sell"), [_m1(1, 100, 101.2, 99.9, 101)], 1100), ("SL", 101.0))

    def test_bougies_deja_vues_ignorees(self):
        g = dict(self.g(), checked_ts=5)
        self.assertIsNone(m._ghost_resolve(g, [_m1(5, 100, 103, 98, 100)], 1100))

    def test_delai_depasse(self):
        g = self.g()
        res = m._ghost_resolve(g, [_m1(1, 100, 100.5, 99.5, 100.4)], 1000 + m.GHOST_MAX_HOURS * 3600)
        self.assertEqual(res, ("Délai", 100.4))

    def test_pnl_pct_net_de_frais(self):
        self.assertAlmostEqual(m._ghost_pnl_pct(self.g(), 102.0), 2.0 - m.GHOST_FEE_PCT)
        self.assertAlmostEqual(m._ghost_pnl_pct(self.g("sell"), 101.0), -1.0 - m.GHOST_FEE_PCT)

    def test_ouverture_dedoublonnee_et_excel(self):
        sig = dict(SIG, price=100.0, atr_m5=1.0)
        self.assertIsNotNone(m.ghost_open(self.wb, sig, "score<0.75"))
        self.assertIsNone(m.ghost_open(self.wb, sig, "score<0.75"))   # même coin/côté/mode
        ws = self.wb["Fantômes"]
        self.assertEqual(ws.max_row, 2)
        self.assertEqual(ws.cell(2, 8).value, "score<0.75")
        self.assertTrue(os.path.exists(m.GHOST_FILE))

    def test_cycle_complet_tp_alimente_apprentissage(self):
        sig = dict(SIG, price=100.0, atr_m5=1.0)   # SL 98.5 / TP 104
        g = m.ghost_open(self.wb, sig, "régime BEAR")
        later = g["checked_ts"] + 60000
        with mock.patch.object(m, "fetch_candles",
                               lambda coin, iv, n, **k: [_m1(later, 100, 104.2, 99.8, 104)]):
            closed = m.ghost_update(self.wb)
        self.assertEqual(len(closed), 1)
        self.assertEqual(m._load_ghosts(), [])
        ws = self.wb["Fantômes"]
        self.assertEqual(ws.cell(2, 12).value, "TP")
        self.assertAlmostEqual(ws.cell(2, 14).value, 4.0 - m.GHOST_FEE_PCT)
        b = m._load_lessons()["buckets"]["regime_RANGING|buy|MR"]
        self.assertEqual(b["gw"], 1)
        self.assertNotIn("gl", b)

    def test_persistance_redemarrage(self):
        m.ghost_open(self.wb, dict(SIG, price=100.0, atr_m5=1.0), "x")
        m._ghosts = None   # simule un redémarrage
        self.assertEqual(len(m._load_ghosts()), 1)


def _declin(n, last_vol=300.0):
    """Bougies en baisse régulière (chaque bougie rouge), volume fort sur la dernière."""
    cs = []
    for i in range(n):
        px = 120.0 - 0.1 * i
        cs.append({"ts": i, "o": px + 0.05, "h": px + 0.06, "l": px - 0.01, "c": px, "v": 100.0})
    cs[-1]["v"] = last_vol
    return cs

class TestGhostCandidates(BaseTest):
    def scan(self, regime):
        m._ghost_candidates.clear(); m._reject_stats.clear()
        m._lessons_cache.clear(); m._lessons_cache_ts = 0.0
        with mock.patch.object(m, "fetch_candles", lambda coin, iv, n, **k: _declin(n)), \
             mock.patch.object(m, "_current_regime", regime):
            return m.get_signal("ETH")

    def test_mom_bloque_en_ranging_devient_candidat_fantome(self):
        self.assertIsNone(self.scan("RANGING"))
        self.assertEqual(len(m._ghost_candidates), 1)
        c = m._ghost_candidates[0]
        self.assertEqual((c["signal"], c["conditions"]["mode"]), ("sell", "MOM"))
        self.assertEqual(c["ghost_reason"], "MOM bloqué RANGING")
        self.assertEqual(m._reject_stats.get("MOM bloqué RANGING"), 1)

    def test_mom_en_bear_est_un_vrai_signal(self):
        sig = self.scan("BEAR")
        self.assertEqual((sig["signal"], sig["conditions"]["mode"]), ("sell", "MOM"))
        self.assertEqual(m._ghost_candidates, [])


def _mr_sell(iv, n, last_vol5):
    """SELL MR : M15 baissier, M5 en haut de bande avec 2 bougies rouges, M1 neutre."""
    cs = []
    if iv == m.INTERVAL_M15:
        for i in range(n):
            px = 130.0 - 0.2 * i
            cs.append({"ts": i, "o": px + 0.05, "h": px + 0.1, "l": px - 0.1, "c": px, "v": 100.0})
    elif iv == m.INTERVAL_M5:
        for i in range(n - 2):
            px = 100.0 + 0.1 * i
            cs.append({"ts": i, "o": px - 0.05, "h": px + 0.02, "l": px - 0.07, "c": px, "v": 100.0})
        top = cs[-1]["c"]
        cs.append({"ts": n - 2, "o": top + 0.02, "h": top + 0.03, "l": top - 0.03, "c": top - 0.02, "v": 100.0})
        cs.append({"ts": n - 1, "o": top - 0.02, "h": top - 0.01, "l": top - 0.06, "c": top - 0.05,
                   "v": last_vol5})
    else:
        for i in range(n):
            px = 100.0 + (0.05 if i % 2 else -0.05)
            cs.append({"ts": i, "o": 100.0, "h": px + 0.01, "l": px - 0.01, "c": px, "v": 100.0})
    return cs

class TestGhostVolume(BaseTest):
    def scan(self, last_vol5):
        m._ghost_candidates.clear(); m._reject_stats.clear()
        m._lessons_cache.clear(); m._lessons_cache_ts = 0.0
        with mock.patch.object(m, "fetch_candles", lambda coin, iv, n, **k: _mr_sell(iv, n, last_vol5)), \
             mock.patch.object(m, "_current_regime", "RANGING"), \
             mock.patch.object(m, "VOL_MULTIPLIER", 1.2):
            return m.get_signal("ETH")

    def test_mr_rejete_de_justesse_par_volume_devient_fantome_mrv(self):
        self.assertIsNone(self.scan(100.0))   # ratio 1.0 < 1.2
        self.assertEqual(len(m._ghost_candidates), 1)
        c = m._ghost_candidates[0]
        self.assertEqual((c["signal"], c["conditions"]["mode"]), ("sell", "MRv"))
        self.assertEqual(c["ghost_reason"], "MR volume ×1.00<1.2")
        self.assertEqual(m._reject_stats, {"volume M5": 1})

    def test_volume_trop_faible_pas_de_fantome(self):
        self.assertIsNone(self.scan(70.0))    # ratio < GHOST_VOL_MIN
        self.assertEqual(m._ghost_candidates, [])

    def test_volume_suffisant_vrai_signal_mr(self):
        sig = self.scan(300.0)
        self.assertEqual((sig["signal"], sig["conditions"]["mode"]), ("sell", "MR"))
        self.assertEqual(m._ghost_candidates, [])

    def test_fantome_mrv_clos_sans_toucher_apprentissage(self):
        if os.path.exists(m.GHOST_FILE):
            os.remove(m.GHOST_FILE)
        m._ghosts = None
        wb = m.init_excel()
        self.scan(100.0)
        g = m.ghost_open(wb, m._ghost_candidates[0], m._ghost_candidates[0]["ghost_reason"])
        later = g["checked_ts"] + 60000
        with mock.patch.object(m, "fetch_candles",
                               lambda coin, iv, n, **k: [_m1(later, g["entry"], g["sl"] + 1, g["entry"] - 0.01, g["entry"])]):
            closed = m.ghost_update(wb)
        self.assertEqual([x["status"] for x in closed], ["SL"])
        self.assertEqual(m._load_lessons().get("buckets", {}), {})
        self.assertEqual(wb["Fantômes"].cell(2, 5).value, "MRv")

    def test_heartbeat_par_type(self):
        closed = [{"side": "sell", "mode": "MOM", "pnl_pct": 1.0},
                  {"side": "sell", "mode": "MOM", "pnl_pct": -0.5},
                  {"side": "buy", "mode": "MRv", "pnl_pct": -0.3}]
        self.assertEqual(m._ghost_by_type(closed),
                         "BUY MRv ✓0/1 moy -0.30% · SELL MOM ✓1/2 moy +0.25%")


# ── Boucle principale (smoke test) ───────────────────────────────────────────
class _StopLoop(Exception):
    pass

class LoopInfo(FakeInfo):
    def meta(self): return {"universe": []}

class TestRunLoop(BaseTest):
    def run_loop(self, daily_chg, asset_positions=None, get_signal=None):
        if os.path.exists(m.GHOST_FILE):
            os.remove(m.GHOST_FILE)
        m._ghosts = None
        m._lessons_cache.clear(); m._lessons_cache_ts = 0.0
        if hasattr(m.run, "_regime_cache"):
            del m.run._regime_cache
        info = LoopInfo(
            mids={c: "100" for c in m.WATCHLIST},
            perp={"assetPositions": asset_positions or [],
                  "crossMarginSummary": {"accountValue": "0", "totalMarginUsed": "0"}},
            spot={"balances": [{"coin": "USDC", "total": "111.18", "hold": "0"}]})
        ex = FakeExchange()
        cycles = []
        def fake_sleep(sec):
            if sec >= 5:                       # attente de fin de cycle
                cycles.append(sec)
                if len(cycles) >= 2:
                    raise _StopLoop()
        def fc(coin, iv, n, **k):
            return _daily(daily_chg) if iv == m.INTERVAL_D1 else _declin(n)
        with mock.patch.object(m, "info", info), mock.patch.object(m, "exchange", ex), \
             mock.patch.object(m, "fetch_candles", fc), \
             mock.patch.object(m.time, "sleep", fake_sleep), \
             mock.patch.object(m, "HEARTBEAT_HOURS", 1e-9), \
             mock.patch.object(m, "get_signal", get_signal or m.get_signal):
            with self.assertRaises(_StopLoop):
                m.run()
        return ex, m._load_ghosts()

    def _signal_eth(self):
        """get_signal factice : BUY MR score 0.9 sur ETH (passe tous les filtres en RANGING)."""
        def gs(coin):
            if coin != "ETH":
                return None
            return {"symbol": "ETH", "signal": "buy", "score": 0.9, "atr_m5": 5.0, "price": 100.0,
                    "conditions": {"mode": "MR", "regime": "RANGING", "rsi": 50, "bbp": 0.3,
                                   "vol_ratio": 1.5}}
        return gs

    def test_signal_retenu_ouvre_un_ordre_en_mode_normal(self):
        """Témoin du test suivant : sans GHOST_ONLY le même signal passe un ordre (SL + TP)."""
        with mock.patch.object(m, "GHOST_ONLY", False):
            ex, _ = self.run_loop(-2, get_signal=self._signal_eth())
        self.assertTrue(ex.orders)

    def test_mode_fantome_seul_aucun_ordre(self):
        with mock.patch.object(m, "GHOST_ONLY", True):
            ex, ghosts = self.run_loop(-2, get_signal=self._signal_eth())
        self.assertEqual(ex.orders, [])
        self.assertEqual([(g["symbol"], g["reason"]) for g in ghosts],
                         [("ETH", "fantôme seul (aurait été tradé)")])
        self.assertIn("fantôme seul", dict(self.alerts)["🚀 Bot démarré"])
        hb = [d for t, d in self.alerts if "résumé" in t]
        self.assertTrue(hb and hb[0].startswith("Mode : **fantôme seul**"))

    def test_wallet_partage_position_copytrade_ni_adoptee_ni_rescannee(self):
        """SOL ouvert par BotCopyTrade (absent de bot_state.json) : pas repris au démarrage
        (aucun SL posé, pas de fermeture zombie) et SOL exclu du scan réel."""
        scanned = []
        def gs(coin):
            scanned.append(coin)
            return None
        sol = {"position": {"coin": "SOL", "szi": "-0.5", "entryPx": "100", "unrealizedPnl": "-3"}}
        ex, _ = self.run_loop(-2, asset_positions=[sol], get_signal=gs)
        self.assertEqual(ex.orders, [])
        self.assertIn("⚠️ Position non reprise", self.titles())
        self.assertTrue(scanned)
        self.assertNotIn("SOL", scanned)
        self.assertIn("ETH", scanned)

    def test_bear_scan_fantome_et_heartbeat(self):
        """2 cycles complets de run() en BEAR : aucun ordre réel, SELL MOM suivi en fantôme,
        heartbeat Discord avec la ligne Fantômes."""
        ex, ghosts = self.run_loop(-10)
        self.assertEqual(ex.orders, [])                     # aucun ordre réel en BEAR
        self.assertTrue(ghosts)
        self.assertTrue(all(g["reason"] == "régime BEAR" for g in ghosts))
        self.assertTrue(any(g["mode"] == "MOM" and g["side"] == "sell" for g in ghosts))
        self.assertEqual(len(ghosts), len({(g["symbol"], g["side"], g["mode"]) for g in ghosts}))  # dédoublonné
        hb = [d for t, d in self.alerts if "résumé" in t]
        self.assertTrue(hb and "Fantômes :" in hb[0])

    def test_ranging_mom_bloque_suivi_en_fantome(self):
        ex, ghosts = self.run_loop(-2)
        self.assertEqual(ex.orders, [])
        self.assertTrue(ghosts)
        self.assertTrue(all(g["reason"] == "MOM bloqué RANGING" for g in ghosts))


# ── Régime de marché ─────────────────────────────────────────────────────────
def _daily(chg_pct, chg3_pct=None):
    """8 bougies D1 : variation totale chg_pct, et variation 3j chg3_pct."""
    start, end = 100.0, 100.0 * (1 + chg_pct / 100)
    closes = [start + (end - start) * i / 7 for i in range(8)]
    if chg3_pct is not None:
        closes[4] = end / (1 + chg3_pct / 100)
    return [{"c": c} for c in closes]

class TestRegime(BaseTest):
    def regime(self, btc, eth):
        data = {"BTC": btc, "ETH": eth}
        with mock.patch.object(m, "fetch_candles", lambda coin, iv, n, **k: data[coin]):
            return m.get_market_regime()[0]

    def test_bear_sur_la_moyenne(self):
        # Cas réel du 2026-10-14 : BTC -2.8 %, ETH -7.4 % → moyenne -5.1 % → BEAR
        self.assertEqual(self.regime(_daily(-2.8), _daily(-7.4)), "BEAR")

    def test_ranging(self):
        self.assertEqual(self.regime(_daily(-2.0), _daily(-5.9)), "RANGING")

    def test_bull(self):
        self.assertEqual(self.regime(_daily(6), _daily(7)), "BULL")

    def test_bull_en_pullback_3j_devient_ranging(self):
        self.assertEqual(self.regime(_daily(6, -3), _daily(7, -3)), "RANGING")

    def test_diverging(self):
        self.assertEqual(self.regime(_daily(4), _daily(-4)), "DIVERGING")

    def test_erreur_api_ranging_par_defaut(self):
        def boom(*a, **k): raise RuntimeError("api down")
        with mock.patch.object(m, "fetch_candles", boom):
            self.assertEqual(m.get_market_regime()[0], "RANGING")


# ── Signaux : diagnostic des rejets ──────────────────────────────────────────
def _candles(n, vol_last, vol=100.0, px=100.0):
    cs = [{"ts": i, "o": px, "h": px * 1.001, "l": px * 0.999, "c": px, "v": vol} for i in range(n)]
    cs[-1]["v"] = vol_last
    return cs

class TestSignalDiagnostic(BaseTest):
    def test_volume_bas_compte_comme_rejet_volume_m5(self):
        m._reject_stats.clear()
        def fc(coin, iv, n, **k): return _candles(n, vol_last=10.0)
        with mock.patch.object(m, "fetch_candles", fc):
            self.assertIsNone(m.get_signal("ETH"))
        self.assertEqual(m._reject_stats.get("volume M5"), 1)


# ── Utilitaires ──────────────────────────────────────────────────────────────
class TestUtils(BaseTest):
    def test_round_px(self):
        self.assertEqual(m.round_px(2012.345), 2012.3)
        self.assertEqual(m.round_px(0.08201357), 0.082014)

    def test_size_mult(self):
        self.assertEqual(m._size_mult(1.6), 1.15)
        self.assertEqual(m._size_mult(1.0), 1.0)
        self.assertEqual(m._size_mult(0.5), 0.85)

    def test_rsi_bornes(self):
        self.assertEqual(m._rsi([float(i) for i in range(30)], 14), 100.0)
        self.assertLess(m._rsi([float(30 - i) for i in range(30)], 14), 1.0)

    def test_bb_plat(self):
        self.assertEqual(m._bb([100.0] * 25)[3], 0.5)


class TestDiscordJournal(unittest.TestCase):
    """Chaque message Discord est copié dans EVENTS_FILE (relu par tools/status.py), même sans webhook."""
    def setUp(self):
        if os.path.exists(m.EVENTS_FILE):
            os.remove(m.EVENTS_FILE)

    def test_message_journalise_sans_webhook(self):
        with mock.patch.object(m, "DISCORD_WEBHOOK_URL", ""):
            m.send_discord_alert("💓 Bot actif", "ligne 1\nligne 2")
        with open(m.EVENTS_FILE, encoding="utf-8") as f:
            ev = [json.loads(l) for l in f]
        self.assertEqual(len(ev), 1)
        self.assertEqual((ev[0]["title"], ev[0]["description"]), ("💓 Bot actif", "ligne 1\nligne 2"))

    def test_journal_en_erreur_n_empeche_pas_l_envoi(self):
        sent = []
        with mock.patch.object(m, "EVENTS_FILE", os.path.join(_TMP, "absent", "x.jsonl")), \
             mock.patch.object(m, "DISCORD_WEBHOOK_URL", "http://hook"), \
             mock.patch.object(m.requests, "post", lambda *a, **k: sent.append(k) or mock.Mock()), \
             mock.patch("sys.stdout", new_callable=io.StringIO):
            m.send_discord_alert("t", "d")
        self.assertEqual(len(sent), 1)


if __name__ == "__main__":
    unittest.main()
