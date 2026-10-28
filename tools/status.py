"""
État des deux bots en une commande, sans copier-coller Discord :
  python3 tools/status.py [--hours 24] [--since "2026-10-27 17:09"]

Lit uniquement des fichiers locaux et le journal systemd (aucun secret, aucun appel réseau) :
  - BotShortTrade : service systemd, erreurs du journal, messages Discord (discord_events.jsonl),
    fantômes clos dans trades.xlsx (onglet Fantômes)
  - BotCopyTrade  : processus, copies simulées ouvertes, messages Discord, dernier état de bot.log
"""
from __future__ import annotations
import argparse, json, os, re, subprocess
from datetime import datetime, timedelta, timezone

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
COPY_ROOT = os.path.join(os.path.dirname(ROOT), "BotCopyTrade")
GHOST_ONLY_SINCE = "2026-10-27 16:09"   # UTC — début du mode fantôme seul
ANSI = re.compile(r"\x1b\[[0-9;]*m")


def sh(cmd: list[str]) -> str:
    try:
        return subprocess.run(cmd, capture_output=True, text=True, timeout=30).stdout
    except Exception as e:
        return f"({e})"


def read_events(path: str, since: datetime) -> list[dict]:
    """Messages Discord journalisés depuis `since` (lignes JSON illisibles ignorées)."""
    out = []
    if not os.path.exists(path):
        return out
    with open(path, encoding="utf-8") as f:
        for line in f:
            try:
                ev = json.loads(line)
                if datetime.fromisoformat(ev["t"]) >= since:
                    out.append(ev)
            except Exception:
                continue
    return out


def print_events(events: list[dict], full_prefix: str):
    """Un message par ligne ; le dernier dont le titre commence par `full_prefix` en entier."""
    if not events:
        print("  (aucun message — le journal local n'existe que depuis le redémarrage qui l'a ajouté)")
        return
    for ev in events:
        first = (ev.get("description") or "").split("\n")[0]
        print(f"  {ev['t'][5:16]}  {ev['title']} — {first[:110]}")
    last = [ev for ev in events if ev["title"].startswith(full_prefix)]
    if last:
        ev = last[-1]
        print(f"\n  Dernier « {ev['title']} » ({ev['t'][5:16]}) :")
        for line in (ev.get("description") or "").split("\n"):
            print(f"    {line}")
        for k, v in (ev.get("fields") or {}).items():
            print(f"    {k} : {v}")


def stats(pnls: list[float]) -> str:
    n = len(pnls)
    mu = sum(pnls) / n
    sd = (sum((p - mu) ** 2 for p in pnls) / (n - 1)) ** 0.5 if n > 1 else 0.0
    t = mu / (sd / n ** 0.5) if sd > 0 else 0.0
    return f"n={n:4d}  WR={100 * sum(p > 0 for p in pnls) / n:3.0f}%  moy={mu:+.3f}%  t={t:+.1f}"


def ghost_report(since: str):
    path = os.path.join(ROOT, "trades.xlsx")
    if not os.path.exists(path):
        print("  trades.xlsx absent"); return
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    if "Fantômes" not in wb.sheetnames:
        print("  pas d'onglet Fantômes"); return
    rows = list(wb["Fantômes"].iter_rows(values_only=True))
    h = {name: i for i, name in enumerate(rows[0])}
    closed, n_open = [], 0
    for r in rows[1:]:
        if r[h["Statut"]] == "Ouvert":
            n_open += 1
        elif r[h["P&L %"]] is not None and (r[h["Fermé (UTC)"]] or "") >= since:
            closed.append(r)
    print(f"  Fantômes clos depuis {since} UTC : {len(closed)} · encore ouverts : {n_open}")
    if not closed:
        return
    groups = {"aurait été tradé (stratégie actuelle)": [r for r in closed if "aurait été tradé" in (r[h["Raison"]] or "")],
              "filtrés / bloqués": [r for r in closed if "aurait été tradé" not in (r[h["Raison"]] or "")]}
    for label, rs in groups.items():
        if not rs:
            continue
        print(f"\n  {label} :  {stats([r[h['P&L %']] for r in rs])}")
        by: dict[str, list[float]] = {}
        for r in rs:
            by.setdefault(f"{r[h['Côté']]} {r[h['Mode']]}", []).append(r[h["P&L %"]])
        for k, v in sorted(by.items()):
            print(f"    {k:10s} {stats(v)}")


def copy_process_running() -> bool:
    for pid in filter(str.isdigit, os.listdir("/proc")):
        try:
            if os.readlink(f"/proc/{pid}/cwd") == COPY_ROOT and \
                    b"main.py" in open(f"/proc/{pid}/cmdline", "rb").read():
                return True
        except OSError:
            continue
    return False


def copy_report(since: datetime):
    print(f"  Processus : {'actif' if copy_process_running() else 'ARRÊTÉ'}")
    path = os.path.join(COPY_ROOT, "simulation_copies.json")
    if os.path.exists(path):
        copies = json.load(open(path, encoding="utf-8"))
        print(f"  Copies simulées ouvertes : {len(copies)}")
        for c in copies.values():
            age = (datetime.now().timestamp() - c["opened_at"]) / 3600
            print(f"    {c['trader_name']:14s} {c['coin']:6s} {'LONG' if c['is_buy'] else 'SHORT':5s} "
                  f"entrée {c['entry']:g}  SL {c['sl']:g}  {age:.0f} h")
    log = os.path.join(COPY_ROOT, "bot.log")
    if os.path.exists(log):
        with open(log, "rb") as f:
            f.seek(max(0, os.path.getsize(log) - 20000))
            tail = ANSI.sub("", f.read().decode("utf-8", "replace"))
        total = [l.strip() for l in tail.splitlines() if "Total P&L copies" in l]
        if total:
            print(f"  bot.log : {total[-1]}")
    print("\n  Messages Discord :")
    events = read_events(os.path.join(COPY_ROOT, "discord_events.jsonl"), since)
    closes = [e for e in read_events(os.path.join(COPY_ROOT, "discord_events.jsonl"),
                                     datetime(2000, 1, 1, tzinfo=timezone.utc))
              if e["title"].startswith("⚪ Fermeture")]
    print_events(events, "📊")
    if closes:
        pnls = [float(e["fields"]["P&L"].replace("$", "")) for e in closes]
        print(f"\n  Copies fermées (total) : {len(closes)} · gagnantes {sum(p > 0 for p in pnls)} "
              f"· P&L ${sum(pnls):+.2f}")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--hours", type=float, default=24, help="fenêtre du journal et des messages (h)")
    ap.add_argument("--since", default=GHOST_ONLY_SINCE, help="date UTC de début du bilan des fantômes")
    a = ap.parse_args()
    since = datetime.now(timezone.utc) - timedelta(hours=a.hours)

    print(f"═══ BotShortTrade ═══ (dernières {a.hours:g} h)")
    state = sh(["systemctl", "--user", "show", "botshorttrade", "-p", "ActiveState,ActiveEnterTimestamp"])
    print("  Service : " + " · ".join(state.strip().splitlines()))
    journal = ANSI.sub("", sh(["journalctl", "--user", "-u", "botshorttrade", "--since", f"-{int(a.hours * 60)}min",
                               "--no-pager", "-o", "cat"]))
    errs: dict[str, int] = {}
    for line in journal.splitlines():
        if re.search(r"ERR|WARN|Traceback|Exception", line):
            key = re.sub(r"\d\d:\d\d:\d\d", "", line.strip())
            errs[key] = errs.get(key, 0) + 1
    print(f"  Journal : {len(journal.splitlines())} lignes · {sum(errs.values())} erreurs/alertes")
    for k, n in sorted(errs.items(), key=lambda kv: -kv[1])[:10]:
        print(f"    ×{n:<3d} {k[:140]}")
    print("\n  Messages Discord :")
    print_events(read_events(os.path.join(ROOT, "discord_events.jsonl"), since), "💓")
    print()
    ghost_report(a.since)

    print(f"\n═══ BotCopyTrade ═══")
    copy_report(since)


if __name__ == "__main__":
    main()
